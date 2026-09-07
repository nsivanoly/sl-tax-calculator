"""Excel generation for tax summary — RAMIS-ready .xlsx using openpyxl."""
import io
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, numbers
from openpyxl.utils import get_column_letter

from app.schemas.calculation import TaxBreakdown


_GREEN = "1a7a3a"
_LIGHT_GREEN = "e6f7ed"
_HEADER_FILL = PatternFill(start_color="fafafa", end_color="fafafa", fill_type="solid")
_ACCENT_FILL = PatternFill(start_color=_LIGHT_GREEN, end_color=_LIGHT_GREEN, fill_type="solid")
_RED_FILL = PatternFill(start_color="fff2f0", end_color="fff2f0", fill_type="solid")
_THIN_BORDER = Border(
    left=Side(style="thin", color="d9d9d9"),
    right=Side(style="thin", color="d9d9d9"),
    top=Side(style="thin", color="d9d9d9"),
    bottom=Side(style="thin", color="d9d9d9"),
)
_NUM_FMT = '#,##0.00'


def _f(v: float | Decimal) -> float:
    return float(v)


def generate_tax_summary_xlsx(
    breakdown: TaxBreakdown,
    fiscal_year: str,
    taxpayer_name: str,
    savings: float = 0.0,
    recommended_relief: str = "local",
) -> io.BytesIO:
    """Generate an RAMIS-ready Excel tax summary and return as a BytesIO buffer."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Tax Summary"

    # Column widths
    ws.column_dimensions["A"].width = 35
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 22

    row = 1

    # --- Title ---
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(row=row, column=1, value="Sri Lanka Income Tax Summary")
    cell.font = Font(size=16, bold=True, color=_GREEN)
    row += 1

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    ws.cell(row=row, column=1, value=f"Fiscal Year {fiscal_year} — {taxpayer_name}").font = Font(size=11, color="808080")
    row += 1

    row += 1

    def _section_header(title: str) -> int:
        nonlocal row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(size=12, bold=True, color=_GREEN)
        cell.fill = _ACCENT_FILL
        for c in range(1, 5):
            ws.cell(row=row, column=c).fill = _ACCENT_FILL
            ws.cell(row=row, column=c).border = _THIN_BORDER
        row += 1
        return row

    def _data_row(label: str, value: float, bold: bool = False, col_b: bool = True) -> int:
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = Font(bold=bold)
        ws.cell(row=row, column=1).border = _THIN_BORDER
        col = 2 if col_b else 4
        val_cell = ws.cell(row=row, column=col, value=value)
        val_cell.number_format = _NUM_FMT
        val_cell.font = Font(bold=bold)
        val_cell.alignment = Alignment(horizontal="right")
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = _THIN_BORDER
        row += 1
        return row

    def _text_row(label: str, value: str, bold: bool = False) -> int:
        nonlocal row
        ws.cell(row=row, column=1, value=label).font = Font(bold=bold)
        ws.cell(row=row, column=1).border = _THIN_BORDER
        ws.cell(row=row, column=2, value=value).font = Font(bold=bold)
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = _THIN_BORDER
        row += 1
        return row

    # --- 1. Gross Income ---
    _section_header("1. Gross Income")
    _data_row("Salary", _f(breakdown.gross_income.salary))
    _data_row("Interest", _f(breakdown.gross_income.interest))
    _data_row("Foreign Currency Income", _f(breakdown.gross_income.foreign_employment))
    _data_row("Other", _f(breakdown.gross_income.other))
    _data_row("Total Gross Income", _f(breakdown.gross_income.total), bold=True)
    row += 1

    # --- 2. Exemptions & Relief ---
    _section_header("2. Exemptions & Relief")
    _data_row("Interest Exemption", _f(breakdown.exemptions.interest_exempt_amount))
    _data_row("Tax-Free Allowance", _f(breakdown.tax_free_allowance))
    row += 1

    # --- 3. Domestic Tax Slab Breakdown ---
    _section_header("3. Domestic Income Tax")
    _data_row("Domestic Income (after exemptions)", _f(breakdown.domestic_tax.domestic_income))
    _data_row("Relief Applied", _f(breakdown.domestic_tax.relief_applied))
    _data_row("Taxable Income", _f(breakdown.domestic_tax.taxable_income), bold=True)
    row += 1

    # Slab table header
    for c, h in enumerate(["Slab", "Taxable Amount", "Rate", "Tax"], 1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = Font(bold=True, size=9)
        cell.fill = _HEADER_FILL
        cell.border = _THIN_BORDER
        cell.alignment = Alignment(horizontal="right" if c > 1 else "left")
    row += 1

    for slab in breakdown.domestic_tax.slab_breakdown:
        ws.cell(row=row, column=1, value=slab.label).border = _THIN_BORDER
        amt_cell = ws.cell(row=row, column=2, value=_f(slab.taxable_in_slab))
        amt_cell.number_format = _NUM_FMT
        amt_cell.alignment = Alignment(horizontal="right")
        amt_cell.border = _THIN_BORDER
        rate_cell = ws.cell(row=row, column=3, value=f"{_f(slab.rate) * 100:.0f}%")
        rate_cell.alignment = Alignment(horizontal="right")
        rate_cell.border = _THIN_BORDER
        tax_cell = ws.cell(row=row, column=4, value=_f(slab.tax))
        tax_cell.number_format = _NUM_FMT
        tax_cell.alignment = Alignment(horizontal="right")
        tax_cell.border = _THIN_BORDER
        row += 1

    # Domestic tax total
    ws.cell(row=row, column=1, value="Total Domestic Tax").font = Font(bold=True)
    ws.cell(row=row, column=1).border = _THIN_BORDER
    total_cell = ws.cell(row=row, column=4, value=_f(breakdown.domestic_tax.tax))
    total_cell.number_format = _NUM_FMT
    total_cell.font = Font(bold=True)
    total_cell.alignment = Alignment(horizontal="right")
    for c in range(1, 5):
        ws.cell(row=row, column=c).border = _THIN_BORDER
        ws.cell(row=row, column=c).fill = PatternFill(start_color="fff7e6", end_color="fff7e6", fill_type="solid")
    row += 2

    # --- 4. Foreign Tax Slab Breakdown ---
    if breakdown.foreign_tax.slab_breakdown:
        _section_header("4. Foreign Income Tax")
        _data_row("Foreign Income", _f(breakdown.foreign_tax.foreign_income))
        _data_row("Relief Applied", _f(breakdown.foreign_tax.relief_applied))
        _data_row("Taxable Income", _f(breakdown.foreign_tax.taxable_income), bold=True)
        row += 1

        for c, h in enumerate(["Slab", "Taxable Amount", "Rate", "Tax"], 1):
            cell = ws.cell(row=row, column=c, value=h)
            cell.font = Font(bold=True, size=9)
            cell.fill = _HEADER_FILL
            cell.border = _THIN_BORDER
            cell.alignment = Alignment(horizontal="right" if c > 1 else "left")
        row += 1

        for slab in breakdown.foreign_tax.slab_breakdown:
            ws.cell(row=row, column=1, value=slab.label).border = _THIN_BORDER
            amt_cell = ws.cell(row=row, column=2, value=_f(slab.taxable_in_slab))
            amt_cell.number_format = _NUM_FMT
            amt_cell.alignment = Alignment(horizontal="right")
            amt_cell.border = _THIN_BORDER
            rate_cell = ws.cell(row=row, column=3, value=f"{_f(slab.rate) * 100:.0f}%")
            rate_cell.alignment = Alignment(horizontal="right")
            rate_cell.border = _THIN_BORDER
            tax_cell = ws.cell(row=row, column=4, value=_f(slab.tax))
            tax_cell.number_format = _NUM_FMT
            tax_cell.alignment = Alignment(horizontal="right")
            tax_cell.border = _THIN_BORDER
            row += 1

        ws.cell(row=row, column=1, value="Total Foreign Tax").font = Font(bold=True)
        ws.cell(row=row, column=1).border = _THIN_BORDER
        total_cell = ws.cell(row=row, column=4, value=_f(breakdown.foreign_tax.tax))
        total_cell.number_format = _NUM_FMT
        total_cell.font = Font(bold=True)
        total_cell.alignment = Alignment(horizontal="right")
        for c in range(1, 5):
            ws.cell(row=row, column=c).border = _THIN_BORDER
            ws.cell(row=row, column=c).fill = PatternFill(start_color="fff7e6", end_color="fff7e6", fill_type="solid")
        row += 2

    # --- 5. Tax Summary ---
    _section_header("5. Tax Summary")
    _data_row("Domestic Tax", _f(breakdown.domestic_tax.tax))
    _data_row("Foreign Tax", _f(breakdown.foreign_tax.tax))
    _data_row("Gross Tax", _f(breakdown.gross_tax), bold=True)
    row += 1

    # --- 6. Credits ---
    _section_header("6. Credits & Deductions")
    _data_row("WHT on Interest", _f(breakdown.credits.wht_on_interest))
    if _f(breakdown.credits.paye_deducted) > 0:
        _data_row("PAYE Deducted", _f(breakdown.credits.paye_deducted))
    if _f(breakdown.credits.self_assessment_paid) > 0:
        _data_row("Self-Assessment Paid", _f(breakdown.credits.self_assessment_paid))
    for adj in breakdown.credits.adjustments:
        amount = -_f(adj.amount) if adj.adjustment_type == "other_deduction" else _f(adj.amount)
        _data_row(adj.label, amount)
    _data_row("Total Credits", _f(breakdown.credits.total_credits), bold=True)
    row += 1

    # --- 7. Final ---
    _section_header("7. Net Tax Payable")
    _data_row("Gross Tax", _f(breakdown.gross_tax))

    # Credits (negative)
    ws.cell(row=row, column=1, value="Less: Total Credits").border = _THIN_BORDER
    cr_cell = ws.cell(row=row, column=2, value=-_f(breakdown.credits.total_credits))
    cr_cell.number_format = _NUM_FMT
    cr_cell.alignment = Alignment(horizontal="right")
    cr_cell.font = Font(color="1890ff")
    for c in range(1, 5):
        ws.cell(row=row, column=c).border = _THIN_BORDER
    row += 1

    # Net payable (big, colored)
    ws.cell(row=row, column=1, value="NET TAX PAYABLE").font = Font(bold=True, size=13)
    ws.cell(row=row, column=1).border = _THIN_BORDER
    net_cell = ws.cell(row=row, column=2, value=_f(breakdown.net_tax_payable))
    net_cell.number_format = _NUM_FMT
    net_cell.font = Font(bold=True, size=13, color="cf1322" if _f(breakdown.net_tax_payable) > 0 else _GREEN)
    net_cell.alignment = Alignment(horizontal="right")
    fill = _RED_FILL if _f(breakdown.net_tax_payable) > 0 else _ACCENT_FILL
    for c in range(1, 5):
        ws.cell(row=row, column=c).border = _THIN_BORDER
        ws.cell(row=row, column=c).fill = fill
    row += 1

    # (end of data rows)

    # Print setup
    ws.print_title_rows = "1:4"
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
